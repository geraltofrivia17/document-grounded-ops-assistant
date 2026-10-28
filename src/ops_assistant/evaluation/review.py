"""Human review sheet for grounding and correctness.

Citation validation only proves that an answer points at real passages that were
supplied. Whether those passages support what the answer says needs a reader, so each
answer run gets a review workbook (``review.xlsx``). The reviewer fills in three
columns per case:

``Correct?``    yes / partly / no: does the answer match the reference answer?
``Supported?``  yes / partly / no / na: is every claim backed by the cited text?
                na when the answer cites nothing (for example a refusal).
``Notes``       why, for every partly or no.

The workbook is laid out for reading: wide wrapped columns, the header and case id kept
in view, and drop-down lists for the marks. Reading the marks back writes
``review_marks.csv`` beside it, so the judgements can be read and diffed without Excel.
Plain CSV review sheets are still accepted.
"""

from __future__ import annotations

import csv
import math
from dataclasses import dataclass
from pathlib import Path

CORRECT_VALUES = ("yes", "partly", "no")
SUPPORTED_VALUES = ("yes", "partly", "no", "na")

# (header, width in characters). What the reviewer reads and fills in comes first, so it
# fits on one screen; the rest follows on the right.
COLUMNS = [
    ("Case", 7),
    ("Question", 30),
    ("Expected answer", 30),
    ("Scoring guidance", 24),
    ("Answer", 44),
    ("Cited passages", 56),
    ("Correct?", 10),
    ("Supported?", 11),
    ("Notes", 28),
    ("Outcome", 15),
    ("Conflict note", 24),
    ("Category", 13),
    ("Type", 15),
    ("Auto check", 9),
]
INPUT_COLUMNS = ("Correct?", "Supported?", "Notes")
CSV_FIELDS = ["case_id", "human_correct", "human_supported", "human_notes"]

INSTRUCTIONS = [
    "How to review",
    "",
    "Read each row from left to right: the question, the expected answer, the scoring guidance, "
    "the answer and the cited passages. Then fill in the three yellow columns. The columns "
    "further right (outcome, conflict note, category, type, automatic check) are for reference.",
    "",
    "Correct?  yes = matches the expected answer and the guidance. partly = the main point "
    "is right but something asked for is missing, or something doubtful is added. "
    "no = wrong, or declines when the handbook answers, or answers when it does not.",
    "For questions the handbook does not cover: yes if it declined without inventing anything.",
    "",
    "Supported?  Judge only from the Cited passages column, not from memory of the handbook. "
    "yes = every claim is stated in the cited text. partly = the main claim is, but at least "
    "one claim is not. no = the cited text does not support the answer. na = nothing cited.",
    "",
    "Notes: say why for every partly and no.",
    "",
    "Save normally (Ctrl+S) and keep the .xlsx format.",
]


@dataclass(frozen=True)
class ReviewMark:
    correct: str
    supported: str
    notes: str


def _row_values(case, record: dict, score) -> list[str]:
    passages = "\n\n".join(
        f"[{chunk_id}] {text}" for chunk_id, text in record.get("cited_text", {}).items()
    )
    values = {
        "Case": case.id,
        "Question": case.question,
        "Expected answer": case.expected_answer,
        "Scoring guidance": case.scoring,
        "Answer": record.get("answer", record.get("error_message", "")),
        "Cited passages": passages,
        "Outcome": record.get("outcome", f"error: {record.get('error_category')}"),
        "Conflict note": record.get("conflict_note", ""),
        "Category": case.category,
        "Type": case.subtype,
        "Auto check": "" if score is None else ("pass" if score.passed else "fail"),
    }
    return [values.get(header, "") for header, _ in COLUMNS]


def _row_height(values: list[str]) -> float:
    """Estimate the height wrapped text needs; Excel won't size rows written by openpyxl."""
    lines = 1
    for value, (_, width) in zip(values, COLUMNS, strict=True):
        text_lines = sum(max(1, math.ceil(len(part) / (width * 1.1))) for part in value.split("\n"))
        lines = max(lines, text_lines)
    return min(409.0, 15.0 * lines + 4)


def write_review_workbook(path: Path, scored: list) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    sheet = book.active
    sheet.title = "Review"
    sheet.append([header for header, _ in COLUMNS])
    for case, record, score in scored:
        sheet.append(_row_values(case, record, score))

    wrap = Alignment(wrap_text=True, vertical="top")
    input_fill = PatternFill("solid", start_color="FFF2CC")
    header_fill = PatternFill("solid", start_color="D9E1F2")
    for index, (header, width) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = width
        for cell in sheet[letter]:
            cell.alignment = wrap
            if header in INPUT_COLUMNS and cell.row > 1:
                cell.fill = input_fill
        sheet[f"{letter}1"].font = Font(bold=True)
        sheet[f"{letter}1"].fill = header_fill
    for row in range(2, sheet.max_row + 1):
        values = [str(sheet.cell(row, c).value or "") for c in range(1, len(COLUMNS) + 1)]
        sheet.row_dimensions[row].height = _row_height(values)
    sheet.freeze_panes = "B2"

    headers = [header for header, _ in COLUMNS]
    last = sheet.max_row
    for header, values in (("Correct?", CORRECT_VALUES), ("Supported?", SUPPORTED_VALUES)):
        letter = get_column_letter(headers.index(header) + 1)
        rule = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        rule.error = f"Choose one of: {', '.join(values)}"
        sheet.add_data_validation(rule)
        rule.add(f"{letter}2:{letter}{last}")

    guide = book.create_sheet("How to review")
    guide.column_dimensions["A"].width = 110
    for line in INSTRUCTIONS:
        guide.append([line])
    for cell in guide["A"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    guide["A1"].font = Font(bold=True, size=13)
    book.save(path)


def _check(case_id: str, correct: str, supported: str) -> None:
    if correct and correct not in CORRECT_VALUES:
        raise ValueError(f"{case_id}: Correct? must be one of {', '.join(CORRECT_VALUES)}")
    if supported and supported not in SUPPORTED_VALUES:
        raise ValueError(f"{case_id}: Supported? must be one of {', '.join(SUPPORTED_VALUES)}")


def _clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def _load_workbook_marks(path: Path) -> dict[str, ReviewMark]:
    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True)["Review"]
    rows = sheet.iter_rows(values_only=True)
    headers = [_clean(h) for h in next(rows)]
    column = {name: headers.index(name) for name in ("Case", *INPUT_COLUMNS)}
    marks: dict[str, ReviewMark] = {}
    for row in rows:
        case_id = _clean(row[column["Case"]])
        if not case_id:
            continue
        correct = _clean(row[column["Correct?"]]).lower()
        supported = _clean(row[column["Supported?"]]).lower()
        _check(case_id, correct, supported)
        marks[case_id] = ReviewMark(correct, supported, _clean(row[column["Notes"]]))
    return marks


def _load_csv_marks(path: Path) -> dict[str, ReviewMark]:
    marks: dict[str, ReviewMark] = {}
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            correct = row["human_correct"].strip().lower()
            supported = row["human_supported"].strip().lower()
            _check(row["case_id"], correct, supported)
            marks[row["case_id"]] = ReviewMark(correct, supported, row["human_notes"].strip())
    return marks


def load_review(path: Path) -> dict[str, ReviewMark]:
    if path.suffix.lower() == ".xlsx":
        return _load_workbook_marks(path)
    return _load_csv_marks(path)


def write_marks_csv(path: Path, marks: dict[str, ReviewMark]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(CSV_FIELDS)
        for case_id, mark in marks.items():
            writer.writerow([case_id, mark.correct, mark.supported, mark.notes])
